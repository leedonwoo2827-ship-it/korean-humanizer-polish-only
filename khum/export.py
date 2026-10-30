"""[반영] → 05_반영본/
  <제목>_윤문본.docx · .txt   채택한 문단은 윤문, 나머지는 원문 그대로
  <제목>_진단표.xlsx          [요약] 등급 · 변경률 · 글자수 · 자가검증 / [문단 대조] / [AI 흔적]
진단만 했으면 진단표만 만든다.
part 를 주면 그 절만: <제목>_<절>_윤문본.docx … (1절을 끝내고 디자인실에 먼저 넘길 때)
"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import audit, config, humanize, taxonomy

HEAD_FILL = PatternFill("solid", fgColor="E9EEF6")
THIN = Side(style="thin", color="C9D1DD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _sheet(ws, cols, rows):
    ws.append([c for c, _ in cols])
    for k, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(k)].width = w
        c = ws.cell(row=1, column=k)
        c.font, c.fill, c.border = Font(bold=True), HEAD_FILL, BOX
    for r in rows:
        ws.append(r)
        for k in range(1, len(cols) + 1):
            c = ws.cell(row=ws.max_row, column=k)
            c.alignment, c.border = WRAP, BOX
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"


def _items_rows(items):
    codes = taxonomy.codes()
    return [[x["para"], x["code"], codes.get(x["code"], {}).get("name", ""), x["severity"], x["old_text"],
             x["reason"], "규칙" if x.get("source") == "rule" else "AI"] for x in items]


ITEM_COLS = [("문단", 6), ("분류", 7), ("패턴", 30), ("심각도", 7), ("원문 구절", 40), ("이유", 50), ("찾은 곳", 7)]


def run(did, decisions=None, part=None):
    d = config.doc(did)
    paras = config.load(did, "01_원문", "paragraphs.json")
    det = config.load(did, "02_탐지", "detect.json")
    if not det:
        raise ValueError("먼저 [AI 흔적 진단]을 실행하세요.")
    rw = config.load(did, "03_윤문", "rewrite.json")
    if decisions is not None:
        config.save(did, "04_검토", "decisions.json", decisions)
    decisions = config.load(did, "04_검토", "decisions.json") or {}
    name = config.slug(d["title"])
    target = d.get("target")
    if part is not None:  # 그 절의 문단만 (목표 글자수는 글 전체 기준이라 절에는 적용하지 않는다)
        pt = humanize.parts(paras)[part]
        lo, hi = pt["start"], pt["end"]
        paras = paras[lo - 1:hi]
        det = dict(det, items=[x for x in det["items"] if lo <= x["para"] <= hi])
        rw = rw and dict(rw, paras=[p for p in rw["paras"] if lo <= p["i"] <= hi])
        name += "_" + config.slug(pt["title"])[:20]
        target = None
    s = audit.summary(paras, det, rw, decisions, target)
    out = config.stage(did, "05_반영본")
    made = []

    if rw:
        fin = audit.final_paras(rw, decisions)
        text = "\n\n".join(t for _, t, _ in fin if t)
        (out / f"{name}_윤문본.txt").write_text(text, encoding="utf-8")
        import docx
        doc = docx.Document()
        for _, t, _ in fin:
            if t:
                doc.add_paragraph(t)
        doc.save(out / f"{name}_윤문본.docx")
        made += [f"{name}_윤문본.docx", f"{name}_윤문본.txt"]

    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    rows = [["글", d["title"]], ["장르", d.get("genre", "")], ["만든 때", datetime.now().strftime("%Y-%m-%d %H:%M")],
            ["찾은 흔적", "S1 {S1} · S2 {S2} · S3 {S3}".format(**s["found"])],
            ["원문 글자수", "공백 포함 {with_space} · 제외 {no_space}".format(**s["orig_chars"])]]
    if rw:
        rows += [["등급", f"{s['grade']} — {s['advice']}"], ["변경률", f"{s['rate']:.0%}"],
                 ["채택 문단", f"{s['adopted']} / {s['n_paras']}"],
                 ["남은 흔적", "S1 {S1} · S2 {S2} · S3 {S3}".format(**s["residual"])],
                 ["윤문본 글자수", "공백 포함 {with_space} · 제외 {no_space}".format(**s["final_chars"])]]
        if s.get("target"):
            t = s["target"]
            rows.append(["목표 글자수", f"{t['chars']}자 ({'공백 제외' if t.get('no_space') else '공백 포함'}) → {t['actual']}자 "
                                    + ("✓ ±5% 안" if t["ok"] else "✗ 벗어남")])
        rows += [[f"자가검증 · {k}", "통과" if v else "확인 필요"] for k, v in s["selfcheck"].items()]
        if s["lost"]:
            rows.append(["빠진 숫자·인용", ", ".join(s["lost"])])
    _sheet(ws, [("항목", 22), ("내용", 90)], rows)

    if rw:
        ws = wb.create_sheet("문단 대조")
        dd = {p["i"]: decisions.get(str(p["i"]), {}) for p in rw["paras"]}
        _sheet(ws, [("문단", 6), ("원문", 60), ("윤문", 60), ("변경률", 8), ("채택", 7), ("비고", 30)],
               [[p["i"], p["orig"], dd[p["i"]].get("text") or p["text"] or "(문단 삭제)", f"{p['rate']:.0%}",
                 "채택" if dd[p["i"]].get("adopt", audit.default_adopt(p)) else "원문 유지",
                 " / ".join(filter(None, ["직접 고침" if dd[p["i"]].get("text") else "",
                                          "변경률 50% 초과" if p["rate"] > audit.STOP else "",
                                          "변경률 30% 초과" if audit.WARN < p["rate"] <= audit.STOP else "",
                                          ("빠짐: " + ", ".join(p["missing"])) if p["missing"] else ""]))]
                for p in rw["paras"]])
        _sheet(wb.create_sheet("남은 흔적"), ITEM_COLS, _items_rows([x for p in rw["paras"] for x in p.get("residual", [])]))
    _sheet(wb.create_sheet("AI 흔적(원문)"), ITEM_COLS, _items_rows(det["items"]))
    wb.save(out / f"{name}_진단표.xlsx")
    made.append(f"{name}_진단표.xlsx")
    if part is None:
        d["exported"] = datetime.now().strftime("%m-%d %H:%M")
    else:
        d.setdefault("exported_parts", {})[str(part)] = datetime.now().strftime("%m-%d %H:%M")
    config.save_doc(d)
    print(f"[반영] {', '.join(made)} → 05_반영본/")
    return {"files": made, "summary": s}
